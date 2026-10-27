from __future__ import annotations

import csv
import email
import io
import json
import logging
import os
import re
import threading
import time
import xml.etree.ElementTree as ET
from email import policy
from pathlib import Path
from typing import Any

from charset_normalizer import from_bytes

from .config import Settings
from .huggingface_cache import (
    disable_huggingface_cache_symlinks,
    materialize_huggingface_cache_symlinks,
)

log = logging.getLogger(__name__)
_docling_converter: Any | None = None
_DOCLING_CONVERTER_LOCK = threading.Lock()
_DOCLING_CONVERSION_LOCK = threading.Lock()


def _prepare_docling_model_cache() -> None:
    if os.name != "nt":
        return
    from huggingface_hub.constants import HF_HUB_CACHE

    cache_dir = Path(HF_HUB_CACHE)
    materialized = materialize_huggingface_cache_symlinks(cache_dir)
    disable_huggingface_cache_symlinks(cache_dir)
    if materialized:
        log.info(
            "Windows-Hugging-Face-Cache für Docling vorbereitet: %d Symlinks durch Kopien ersetzt",
            materialized,
        )


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".adoc", ".tex", ".log", ".ini",
    ".cfg", ".conf", ".config", ".properties", ".toml", ".yaml", ".yml",
    ".json", ".jsonl", ".xml", ".xsd", ".xsl", ".xslt", ".html", ".htm",
    ".xhtml", ".csv", ".tsv", ".py", ".pyi", ".js", ".jsx", ".ts", ".tsx",
    ".mjs", ".cjs", ".java", ".kt", ".kts", ".cs", ".fs", ".vb", ".go",
    ".rs", ".c", ".h", ".cpp", ".hpp", ".sql", ".ps1", ".psm1", ".bat",
    ".cmd", ".sh", ".bash", ".zsh", ".dockerfile", ".makefile", ".mk", ".r",
    ".ipynb", ".graphql", ".gql", ".proto", ".tf", ".hcl", ".properties",
    ".env.example", ".diff", ".patch", ".vbs", ".php", ".rb", ".swift", ".pl",
    ".lua", ".dart", ".ex", ".exs", ".erl", ".hrl", ".clj", ".cljs", ".scala",
    ".jenkinsfile", ".procfile", ".gitignore", ".dockerignore", ".gitattributes",
    ".gitmodules", ".editorconfig", ".npmrc", ".nvmrc", ".prettierrc",
    ".python-version", ".eslintrc", ".flake8", ".coveragerc",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".eml", ".msg"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS | {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".rtf"}


def _decode(raw: bytes) -> str:
    if not raw:
        return ""
    try:
        match = from_bytes(raw).best()
        if match is not None:
            return str(match)
    except Exception:
        pass
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_csv(text: str, delimiter: str | None = None) -> str:
    sample = text[:8192]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows: list[str] = []
    for row_number, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        cells = [str(cell).strip() for cell in row]
        if any(cells):
            rows.append(f"Zeile {row_number}: " + " | ".join(cells))
    return "\n".join(rows)


def _extract_docx(raw: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(raw))
    output: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = getattr(paragraph.style, "name", "") or ""
        if style.lower().startswith("heading") or style.lower().startswith("überschrift"):
            output.append(f"\n## {value}\n")
        else:
            output.append(value)
    for table_number, table in enumerate(document.tables, start=1):
        output.append(f"\nTabelle {table_number}:")
        for row in table.rows:
            cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            output.append(" | ".join(cells))
    return "\n".join(output)


def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    output: list[str] = []
    try:
        for sheet in workbook.worksheets:
            output.append(f"\n## Arbeitsblatt: {sheet.title}\n")
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = ["" if value is None else str(value).strip() for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    output.append(f"Zeile {row_number}: " + " | ".join(cells))
    finally:
        workbook.close()
    return "\n".join(output)


def _extract_pptx(raw: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(raw))
    output: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        output.append(f"\n## Folie {slide_number}\n")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    output.append(value)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    output.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(output)


def _extract_pdf(raw: bytes) -> str:
    import fitz

    output: list[str] = []
    with fitz.open(stream=raw, filetype="pdf") as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                output.append(f"\n## Page {page_number}\n{text}")
    return "\n".join(output)


def _extract_eml(raw: bytes) -> str:
    message = email.message_from_bytes(raw, policy=policy.default)
    output = [
        f"Betreff: {message.get('subject', '')}",
        f"Von: {message.get('from', '')}",
        f"An: {message.get('to', '')}",
        f"Datum: {message.get('date', '')}",
    ]
    body_parts: list[str] = []
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_type() == "text/plain" and part.get_content_disposition() != "attachment":
                try:
                    body_parts.append(part.get_content())
                except (LookupError, UnicodeDecodeError):
                    payload = part.get_payload(decode=True) or b""
                    body_parts.append(_decode(payload))
    elif message.get_content_type() == "text/plain":
        try:
            body_parts.append(message.get_content())
        except Exception:
            body_parts.append(_decode(message.get_payload(decode=True) or b""))
    output.extend(body_parts)
    return "\n".join(str(item) for item in output if item)


def _extract_html(text: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    for element in soup(["script", "style", "noscript", "svg"]):
        element.decompose()
    return soup.get_text("\n", strip=True)


def _extract_xml(text: str) -> str:
    try:
        root = ET.fromstring(text)
        lines: list[str] = []
        for element in root.iter():
            value = (element.text or "").strip()
            if value:
                tag = element.tag.rsplit("}", 1)[-1]
                lines.append(f"{tag}: {value}")
        return "\n".join(lines)
    except ET.ParseError:
        return text


def _extract_ipynb(text: str) -> str:
    try:
        notebook = json.loads(text)
        output: list[str] = []
        for number, cell in enumerate(notebook.get("cells", []), start=1):
            kind = cell.get("cell_type", "cell")
            source = cell.get("source", "")
            if isinstance(source, list):
                source = "".join(source)
            if source:
                output.append(f"\n## Notebook {kind} {number}\n{source}")
            for out in cell.get("outputs", []):
                data = out.get("data", {})
                plain = data.get("text/plain", "")
                if isinstance(plain, list):
                    plain = "".join(plain)
                if plain:
                    output.append(str(plain))
        return "\n".join(output)
    except (ValueError, TypeError):
        return text


def _docling_extract(path: Path) -> str | None:
    global _docling_converter
    try:
        from docling.document_converter import DocumentConverter
    except ImportError as exc:
        raise RuntimeError(
            "USE_DOCLING=true, aber Docling ist in dieser Installation nicht verfügbar. "
            "Installiere die Docling-Abhängigkeit mit `pip install -e '.[docling]'` "
            "oder verwende einen Release-Build mit Docling."
        ) from exc
    try:
        if _docling_converter is None:
            with _DOCLING_CONVERTER_LOCK:
                if _docling_converter is None:
                    _prepare_docling_model_cache()
                    started = time.monotonic()
                    _docling_converter = DocumentConverter()
                    log.info(
                        "Docling-Pipeline initialisiert (%.1f s); verwende sie für weitere Dokumente wieder.",
                        time.monotonic() - started,
                    )
        started = time.monotonic()
        with _DOCLING_CONVERSION_LOCK:
            result = _docling_converter.convert(str(path))
        log.info("Docling-Konvertierung %s: %.1f s", path.name, time.monotonic() - started)
        return result.document.export_to_markdown()
    except Exception as exc:
        log.warning("Docling konnte %s nicht extrahieren: %s; verwende Standardparser.", path, exc)
        return None


def extract_text(path: Path, raw: bytes, settings: Settings) -> str:
    """Extract text from common document and source-code formats; return '' if unsupported."""
    from .scanner import extension_for_path
    suffix = extension_for_path(path)
    if suffix not in SUPPORTED_EXTENSIONS:
        return ""

    if settings.use_docling and suffix in {
        ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".doc", ".xls", ".ppt",
        ".odt", ".ods", ".odp", ".rtf", ".html", ".htm", ".md", ".csv",
        ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".msg",
    }:
        extracted = _docling_extract(path)
        if extracted and extracted.strip():
            return extracted.strip()

    if suffix == ".pdf":
        return _extract_pdf(raw)
    if suffix in {".docx"}:
        return _extract_docx(raw)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(raw)
    if suffix == ".pptx":
        return _extract_pptx(raw)
    if suffix in IMAGE_EXTENSIONS:
        # Images are metadata-only unless the optional Docling/OCR pipeline is enabled.
        return ""
    if suffix in {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".rtf", ".msg"}:
        # These legacy formats are handled through Docling when enabled. Avoid
        # interpreting their binary containers as text when Docling is not installed.
        return ""

    text = _decode(raw)
    if "\x00" in text[:2048] and suffix not in {".ipynb"}:
        return ""
    if suffix in {".html", ".htm", ".xhtml"}:
        return _extract_html(text)
    if suffix in {".csv", ".tsv"}:
        return _extract_csv(text, "\t" if suffix == ".tsv" else None)
    if suffix in {".json"}:
        try:
            return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except (ValueError, TypeError):
            return text
    if suffix == ".jsonl":
        lines = []
        for line in text.splitlines():
            try:
                lines.append(json.dumps(json.loads(line), ensure_ascii=False))
            except (ValueError, TypeError):
                lines.append(line)
        return "\n".join(lines)
    if suffix in {".yaml", ".yml"}:
        try:
            import yaml
            value: Any = yaml.safe_load(text)
            return yaml.safe_dump(value, allow_unicode=True, sort_keys=False)
        except Exception:
            return text
    if suffix in {".xml", ".xsd", ".xsl", ".xslt"}:
        return _extract_xml(text)
    if suffix == ".ipynb":
        return _extract_ipynb(text)
    return re.sub(r"\x00+", "", text).strip()
